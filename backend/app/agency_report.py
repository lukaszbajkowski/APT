"""Comiesięczne zestawienie osób dla agencji pracy (APT).

Do zestawienia trafia każdy, kto choć jeden dzień miesiąca był zatrudniony i przypisany do danej APT.
Okresy liczymy z:
  - bieżącego zatrudnienia (data zatrudnienia – data zwolnienia z kartoteki),
  - wcześniejszych okresów (ponowne zatrudnienie – tabela employment_periods),
  - historii przydziałów (zmiana APT w trakcie zatrudnienia – tabela assignments).
"""

import calendar
import io
from dataclasses import dataclass
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Agency, Assignment, Employee, EmploymentPeriod

MONTHS = [
    "",
    "styczeń",
    "luty",
    "marzec",
    "kwiecień",
    "maj",
    "czerwiec",
    "lipiec",
    "sierpień",
    "wrzesień",
    "październik",
    "listopad",
    "grudzień",
]
FAR = date(9999, 12, 31)


@dataclass
class Row:
    employee_id: int
    number: str
    last_name: str
    first_name: str
    department: str | None
    position: str | None
    hire_date: date | None
    termination_date: date | None
    period_from: date  # w obrębie miesiąca
    period_to: date
    days: int
    change: str  # przyjęty | odszedł | przyjęty i odszedł | ""


def month_range(month: str) -> tuple[date, date]:
    """„2026-10” → (1.10.2026, 31.10.2026)."""
    try:
        y, m = (int(x) for x in month.split("-"))
        start = date(y, m, 1)
    except (ValueError, TypeError) as exc:
        raise ValueError("Miesiąc w formacie RRRR-MM.") from exc
    return start, date(y, m, calendar.monthrange(y, m)[1])


def month_label(month: str) -> str:
    start, _ = month_range(month)
    return f"{MONTHS[start.month]} {start.year}"


def previous_month(today: date | None = None) -> str:
    first = (today or date.today()).replace(day=1)
    prev = first - timedelta(days=1)
    return f"{prev.year}-{prev.month:02d}"


def _agency_segments(
    assignments: list[Assignment], start: date | None, end: date | None, current: str | None
):
    """Odcinki (od, do, nazwa APT) w okresie zatrudnienia na podstawie historii przydziałów."""
    start = start or date.min
    end = end or FAR
    hist = sorted(assignments, key=lambda a: (a.valid_from, a.id))
    if not hist:
        return [(start, end, current)]
    out = []
    # przydział obowiązujący w dniu zatrudnienia: ostatni z valid_from <= start (albo pierwszy w ogóle)
    before = [a for a in hist if a.valid_from <= start]
    after = [a for a in hist if a.valid_from > start and a.valid_from <= end]
    first = before[-1] if before else hist[0]
    points = [(start, first.agency_name)] + [(a.valid_from, a.agency_name) for a in after]
    # ostatni odcinek = obecny przydział z kartoteki (historia mogła powstać przed zmianą nazwy APT
    # albo przed wprowadzeniem historii przydziałów)
    points[-1] = (points[-1][0], current)
    for i, (d, name) in enumerate(points):
        to = points[i + 1][0] - timedelta(days=1) if i + 1 < len(points) else end
        if to >= d:
            out.append((d, to, name))
    return out


def rows_for(db: Session, agency: Agency, month: str) -> list[Row]:
    m_start, m_end = month_range(month)
    name = agency.name
    # bez filtra po dacie: ktoś obecnie zatrudniony później mógł być w tej APT w poprzednim okresie
    emps = list(db.scalars(select(Employee)).unique())
    ids = [e.id for e in emps]
    hist: dict[int, list[Assignment]] = {}
    for a in db.scalars(select(Assignment).where(Assignment.employee_id.in_(ids))):
        hist.setdefault(a.employee_id, []).append(a)
    periods: dict[int, list[EmploymentPeriod]] = {}
    for p in db.scalars(select(EmploymentPeriod).where(EmploymentPeriod.employee_id.in_(ids))):
        periods.setdefault(p.employee_id, []).append(p)

    out: list[Row] = []
    for e in emps:
        spans: list[tuple[date, date, date | None, date | None]] = []  # (od, do, zatrudn., zwoln.)
        # bieżący okres
        if e.hire_date is not None:
            current = e.agency.name if e.agency else None
            for d0, d1, ag in _agency_segments(hist.get(e.id, []), e.hire_date, e.termination_date, current):
                if ag == name:
                    spans.append((d0, d1, e.hire_date, e.termination_date))
        # wcześniejsze okresy
        for p in periods.get(e.id, []):
            if p.agency_name == name and p.hire_date:
                spans.append((p.hire_date, p.termination_date or FAR, p.hire_date, p.termination_date))
        for d0, d1, hire, term in spans:
            a, b = max(d0, m_start), min(d1, m_end)
            if a > b:
                continue
            came = hire is not None and m_start <= hire <= m_end
            left = term is not None and m_start <= term <= m_end
            change = (
                "przyjęty i odszedł" if came and left else "przyjęty" if came else "odszedł" if left else ""
            )
            out.append(
                Row(
                    employee_id=e.id,
                    number=f"{e.number:06d}",
                    last_name=e.last_name,
                    first_name=e.first_name,
                    department=e.department.name if e.department else None,
                    position=e.position.name if e.position else None,
                    hire_date=hire,
                    termination_date=term,
                    period_from=a,
                    period_to=b,
                    days=(b - a).days + 1,
                    change=change,
                )
            )
    out.sort(key=lambda r: (r.last_name.lower(), r.first_name.lower(), r.period_from))
    return out


def _d(v: date | None) -> str:
    return v.strftime("%d.%m.%Y") if v else ""


def xlsx(agency: Agency, month: str, rows: list[Row]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = month
    ws["A1"] = f"Zestawienie osób – {agency.name} – {month_label(month)}"
    ws["A1"].font = Font(bold=True, size=13)
    people = len({r.employee_id for r in rows})
    came = sum(1 for r in rows if r.change.startswith("przyjęty"))
    left = sum(1 for r in rows if "odszedł" in r.change)
    ws["A2"] = f"Osób: {people} · przyjętych: {came} · odchodzących: {left}"
    head = [
        "Nr ewid.",
        "Nazwisko",
        "Imię",
        "Wydział",
        "Stanowisko",
        "Data zatrudnienia",
        "Data zwolnienia",
        "Od (w miesiącu)",
        "Do (w miesiącu)",
        "Dni",
        "Zmiana",
    ]
    ws.append([])
    ws.append(head)
    fill = PatternFill("solid", fgColor="1F3A68")
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(
            [
                r.number,
                r.last_name,
                r.first_name,
                r.department or "",
                r.position or "",
                _d(r.hire_date),
                _d(r.termination_date),
                _d(r.period_from),
                _d(r.period_to),
                r.days,
                r.change,
            ]
        )
    widths = [10, 20, 16, 20, 22, 14, 14, 14, 14, 6, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(head))}{max(4, 4 + len(rows))}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_name(agency: Agency, month: str) -> str:
    from app.templating import safe_filename

    return safe_filename(f"Zestawienie {agency.name} {month}") + ".xlsx"


def email(agency: Agency, month: str, rows: list[Row]):  # noqa: ANN201
    from app.emails import Email

    people = len({r.employee_id for r in rows})
    came = sum(1 for r in rows if r.change.startswith("przyjęty"))
    left = sum(1 for r in rows if "odszedł" in r.change)
    return Email(
        subject=f"Zestawienie osób – {agency.name} – {month_label(month)}",
        title="Miesięczne zestawienie osób",
        preheader=f"{people} osób w {month_label(month)}.",
        greeting="Dzień dobry,",
        paragraphs=[
            f"w załączniku przesyłamy zestawienie osób z APT „{agency.name}” za {month_label(month)} "
            "z okresami zatrudnienia w tym miesiącu."
        ],
        details=[("Osób", str(people)), ("Przyjętych", str(came)), ("Odchodzących", str(left))],
        footer_note="Wiadomość automatyczna. W razie rozbieżności prosimy o kontakt z działem kadr.",
        attachments=[
            (
                file_name(agency, month),
                xlsx(agency, month, rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        ],
    )


def recipients(agency: Agency) -> list[str]:
    return [x.strip() for x in (agency.email or "").split(",") if x.strip()]


def send(
    db: Session, agency: Agency, month: str, *, sync: bool = False, user=None, request=None
) -> list[str]:  # noqa: ANN001
    """Wysyła zestawienie na adresy agencji (+ kopia do kadr). Zwraca listę adresatów."""
    from app import mailer
    from app.audit import audit
    from app.models import Role, User

    rows = rows_for(db, agency, month)
    mail = email(agency, month, rows)
    to = recipients(agency)
    hr = [u.email for u in db.scalars(select(User).where(User.role == Role.hr, User.is_active)) if u.email]
    sent: list[str] = []
    for addr in to + [h for h in hr if h not in to]:
        if sync and mailer.outbox is None:
            try:
                mailer.send_now(addr, mail)
                sent.append(addr)
            except Exception as exc:  # noqa: BLE001
                print(f"Błąd wysyłki zestawienia do {addr}: {exc}")
        elif mailer.send(addr, mail):
            sent.append(addr)
    audit(
        db,
        request,
        "agency_report_sent",
        user=user,
        username=None if user else "system",
        entity_type="agency",
        entity_id=agency.id,
        details={
            "agency": agency.name,
            "month": month,
            "people": len({r.employee_id for r in rows}),
            "to": sent,
        },
    )
    db.commit()
    return sent


def send_monthly(db: Session, today: date | None = None) -> int:
    """1. dnia miesiąca: zestawienia za poprzedni miesiąc dla APT z adresem e-mail (raz na miesiąc)."""
    from app.models import AuditLog

    month = previous_month(today)
    done = {
        (a.details or {}).get("agency")
        for a in db.scalars(select(AuditLog).where(AuditLog.action == "agency_report_sent"))
        if (a.details or {}).get("month") == month and a.username == "system"
    }
    n = 0
    for ag in db.scalars(select(Agency).where(Agency.is_active, Agency.email.is_not(None))):
        if ag.name in done or not recipients(ag):
            continue
        send(db, ag, month, sync=True)
        n += 1
    return n
