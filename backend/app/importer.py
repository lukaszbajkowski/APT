"""Import pracowników z Excela (.xlsx) i CSV: odczyt pliku, dopasowanie kolumn, normalizacja wartości.

Ten moduł nie dotyka bazy – walidacja względem słowników i duplikatów jest w routers/imports.py.
"""

import csv
import io
import re
import tempfile
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from app.pesel import PeselError, parse

MAX_ROWS = 5000


class ImportFileError(ValueError):
    pass


@dataclass
class Table:
    headers: list[str]
    rows: list[list[Any]]  # bez wiersza nagłówka i pustych wierszy
    row_numbers: list[
        int
    ]  # numery wierszy w pliku (1 = pierwszy), żeby komunikaty wskazywały wiersz z Excela
    sheet: str | None = None


@dataclass(frozen=True)
class FieldSpec:
    key: str
    label: str
    synonyms: tuple[str, ...]
    required: bool = False


# Kolejność = kolejność w kreatorze i w szablonie.
FIELDS: tuple[FieldSpec, ...] = (
    FieldSpec("last_name", "Nazwisko", ("nazwisko", "lastname", "surname", "nazwiskopracownika")),
    FieldSpec("first_name", "Imię", ("imie", "imiona", "firstname", "name", "imiepracownika")),
    FieldSpec(
        "full_name",
        "Nazwisko i imię (w jednej kolumnie)",
        ("nazwiskoiimie", "nazwiskoimie", "imieinazwisko", "pracownik", "osoba", "fullname"),
    ),
    FieldSpec("pesel", "PESEL", ("pesel", "nrpesel", "numerpesel")),
    FieldSpec(
        "birth_date", "Data urodzenia", ("dataurodzenia", "urodzony", "urodzona", "dataur", "birthdate")
    ),
    FieldSpec("sex", "Płeć", ("plec", "kobietamezczyzna", "km", "sex", "gender")),
    FieldSpec(
        "country",
        "Obywatelstwo",
        ("obywatelstwo", "narodowosc", "kraj", "krajpochodzenia", "nationality", "citizenship"),
    ),
    FieldSpec("agency", "APT", ("apt", "agencja", "agencjapracy", "agencjapracytymczasowej", "agency")),
    FieldSpec("department", "Wydział", ("wydzial", "dzial", "department", "oddzial")),
    FieldSpec("position", "Stanowisko", ("stanowisko", "position", "funkcja")),
    FieldSpec(
        "hire_date",
        "Data zatrudnienia",
        (
            "datazatrudnienia",
            "zatrudniony",
            "zatrudnionaod",
            "zatrudnionyod",
            "dataprzyjecia",
            "od",
            "hiredate",
        ),
    ),
    FieldSpec("notes", "Uwagi", ("uwagi", "notatki", "komentarz", "notes")),
)
FIELD_KEYS = {f.key for f in FIELDS}


def norm_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).replace("ł", "l").replace("Ł", "L")
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def suggest_mapping(headers: list[str], fields: tuple["FieldSpec", ...] | None = None) -> dict[str, int]:
    """Proponuje przypisanie: pole -> indeks kolumny. Każda kolumna co najwyżej raz."""
    normalized = [norm_header(h) for h in headers]
    mapping: dict[str, int] = {}
    used: set[int] = set()
    # 1. dokładne dopasowania, 2. nagłówek zaczynający się od synonimu (np. "Data urodzenia (dd.mm.rrrr)")
    for exact in (True, False):
        for spec in fields or FIELDS:
            if spec.key in mapping:
                continue
            for i, h in enumerate(normalized):
                if i in used or not h:
                    continue
                hit = (
                    h in spec.synonyms
                    if exact
                    else any(len(s) >= 4 and h.startswith(s) for s in spec.synonyms)
                )
                if hit:
                    mapping[spec.key] = i
                    used.add(i)
                    break
    # Jeśli są osobne kolumny nazwiska i imienia, kolumna łączona jest zbędna.
    if "last_name" in mapping and "first_name" in mapping:
        mapping.pop("full_name", None)
    return mapping


# ---------- odczyt pliku ----------


def read_table(raw: bytes, filename: str) -> Table:
    name = (filename or "").lower()
    if name.endswith(".xls"):
        raise ImportFileError("Stary format .xls nie jest obsługiwany. Zapisz plik w Excelu jako .xlsx.")
    if name.endswith(".numbers"):
        return _read_numbers(raw)
    if name.endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK":
        return _read_xlsx(raw)
    if name.endswith((".csv", ".txt")):
        return _read_csv(raw)
    raise ImportFileError("Obsługiwane formaty: .xlsx, .numbers (Numbers na Macu) i .csv.")


def _is_empty(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _from_rows(rows: list[list[Any]], sheet: str | None) -> Table:
    # Nagłówek = pierwszy wiersz z co najmniej dwiema wypełnionymi komórkami.
    header_idx = next(
        (i for i, r in enumerate(rows) if sum(not _is_empty(c) for c in r) >= 2),
        None,
    )
    if header_idx is None:
        raise ImportFileError("Nie znaleziono wiersza nagłówka (nazw kolumn).")
    headers = [str(h).strip() if not _is_empty(h) else "" for h in rows[header_idx]]
    # Obcinamy puste kolumny na końcu.
    while headers and not headers[-1]:
        headers.pop()
    width = len(headers)
    data = []
    numbers = []
    for offset, r in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        r = (list(r) + [None] * width)[:width]
        if all(_is_empty(c) for c in r):
            continue
        data.append(r)
        numbers.append(offset)
    if not data:
        raise ImportFileError("Plik nie zawiera żadnych wierszy z danymi.")
    if len(data) > MAX_ROWS:
        raise ImportFileError(f"Plik ma {len(data)} wierszy. Jednorazowo można zaimportować {MAX_ROWS}.")
    return Table(headers=headers, rows=data, row_numbers=numbers, sheet=sheet)


def _read_xlsx(raw: bytes) -> Table:
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise ImportFileError("Nie udało się otworzyć pliku Excel.") from exc
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if any(sum(not _is_empty(c) for c in r) >= 2 for r in rows):
                return _from_rows(rows, ws.title)
    finally:
        wb.close()
    raise ImportFileError("Arkusz jest pusty.")


def _read_numbers(raw: bytes) -> Table:
    """Arkusz Apple Numbers. Czytamy pierwszą tabelę z danymi (w Numbers arkusz może mieć kilka tabel)."""
    from numbers_parser import Document
    from numbers_parser.exceptions import FileError, FileFormatError

    if raw[:2] != b"PK":
        raise ImportFileError(
            "Nie udało się otworzyć pliku Numbers. W Numbers wybierz Plik → Eksportuj do → Excel "
            "albo zapisz plik ponownie (Plik → Zapisz)."
        )
    with tempfile.NamedTemporaryFile(suffix=".numbers") as tmp:
        tmp.write(raw)
        tmp.flush()
        try:
            doc = Document(tmp.name)
        except (FileError, FileFormatError, KeyError, ValueError, OSError) as exc:
            msg = str(exc).lower()
            if "password" in msg or "encrypted" in msg:
                raise ImportFileError(
                    "Plik Numbers jest chroniony hasłem. Usuń hasło i spróbuj ponownie."
                ) from exc
            raise ImportFileError("Nie udało się otworzyć pliku Numbers.") from exc
        except Exception as exc:  # biblioteka potrafi rzucić różnymi wyjątkami przy nietypowych plikach
            raise ImportFileError("Nie udało się odczytać pliku Numbers.") from exc
        for sheet in doc.sheets:
            for table in sheet.tables:
                rows = [[cell.value for cell in row] for row in table.iter_rows()]
                if any(sum(not _is_empty(c) for c in r) >= 2 for r in rows):
                    name = sheet.name if len(sheet.tables) == 1 else f"{sheet.name} / {table.name}"
                    return _from_rows(rows, name)
    raise ImportFileError("Plik Numbers nie zawiera tabeli z danymi.")


def _read_csv(raw: bytes) -> Table:
    for enc in ("utf-8-sig", "cp1250"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError("Nie rozpoznano kodowania pliku CSV (zapisz jako UTF-8).")
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return _from_rows(rows, None)


# ---------- normalizacja wartości ----------


def cell_text(v: Any) -> str | None:
    if _is_empty(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        v = v.date()
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    return " ".join(str(v).split())


_DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y.%m.%d", "%d.%m.%y")


def parse_date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    text = cell_text(v)
    if text is None:
        raise ValueError("pusta data")
    text = text.split(" ")[0]  # "01.02.1990 00:00"
    for fmt in _DATE_FORMATS:
        try:
            d = datetime.strptime(text, fmt).date()
            if fmt == "%d.%m.%y" and d.year > date.today().year:
                d = d.replace(year=d.year - 100)
            return d
        except ValueError:
            continue
    raise ValueError(f"nierozpoznana data „{text}” (oczekiwano np. 31.12.1990)")


def parse_sex(v: Any) -> str:
    text = (cell_text(v) or "").lower()
    if text in {"k", "kobieta", "f", "female", "ż", "z"}:
        return "K"
    if text in {"m", "mężczyzna", "mezczyzna", "male"}:
        return "M"
    raise ValueError(f"nierozpoznana płeć „{cell_text(v)}” (K/M)")


def parse_pesel(v: Any) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    text = re.sub(r"[\s-]", "", str(v))
    # Excel gubi zero na początku (urodzeni w latach 2000–2009 mają PESEL zaczynający się od 0).
    if text.isdigit() and len(text) == 10:
        text = "0" + text
    parse(text)  # rzuca PeselError
    return text


def split_full_name(v: Any) -> tuple[str, str]:
    """„Kowalski Jan Maria” -> ("Kowalski", "Jan Maria"). Pierwszy wyraz to nazwisko."""
    text = cell_text(v) or ""
    parts = text.split(" ", 1)
    if len(parts) < 2:
        raise ValueError("w kolumnie „nazwisko i imię” brakuje imienia")
    return parts[0], parts[1]


def nice_name(v: str) -> str:
    """WIELKIE LITERY z Excela -> Wielkie Pierwsze; inaczej bez zmian."""
    if v.isupper() and len(v) > 1:
        return " ".join("-".join(p.capitalize() for p in w.split("-")) for w in v.lower().split(" "))
    return v


# Przymiotniki obywatelstwa -> kod ISO (w arkuszach kadrowych często „ukraińskie” zamiast „Ukraina”).
_NATIONALITY = {
    "polsk": "PL",
    "ukrain": "UA",
    "bialorus": "BY",
    "moldaw": "MD",
    "gruzin": "GE",
    "ormian": "AM",
    "armen": "AM",
    "azer": "AZ",
    "kazach": "KZ",
    "kirgi": "KG",
    "uzbec": "UZ",
    "uzbek": "UZ",
    "tadzyk": "TJ",
    "turkmen": "TM",
    "rosyj": "RU",
    "litew": "LT",
    "lotew": "LV",
    "estons": "EE",
    "niemiec": "DE",
    "czes": "CZ",
    "slowac": "SK",
    "rumun": "RO",
    "bulgar": "BG",
    "wegier": "HU",
    "turec": "TR",
    "indyj": "IN",
    "nepal": "NP",
    "bangla": "BD",
    "pakistan": "PK",
    "lankij": "LK",
    "filipin": "PH",
    "indonez": "ID",
    "wietnam": "VN",
    "kolumbij": "CO",
    "wenezuel": "VE",
    "meksyk": "MX",
    "nigeryj": "NG",
    "zimbabw": "ZW",
    "etiop": "ET",
    "egips": "EG",
    "marokan": "MA",
}


def nationality_iso(v: str) -> str | None:
    key = norm_header(v)
    if len(key) == 2:
        return key.upper()
    for stem, iso in _NATIONALITY.items():
        if key.startswith(stem):
            return iso
    return None


# ---------- szablon ----------


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pracownicy"
    cols = [f for f in FIELDS if f.key != "full_name"]
    ws.append([f.label for f in cols])
    ws.append(
        [
            "Kowalski",
            "Jan",
            "",
            "",
            "M",
            "Ukraina",
            "Agencja Przykładowa",
            "Montaż",
            "Operator",
            "01.03.2026",
            "",
        ]
    )
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A73")
    for col, width in zip("ABCDEFGHIJK", (18, 14, 14, 15, 7, 15, 24, 18, 18, 17, 30), strict=True):
        ws.column_dimensions[col].width = width
    # PESEL jako tekst, żeby Excel nie zjadał zer na początku.
    for row in ws.iter_rows(min_row=2, max_row=2000, min_col=3, max_col=3):
        for c in row:
            c.number_format = "@"
    info = wb.create_sheet("Instrukcja")
    for line in (
        "Wypełnij arkusz „Pracownicy” – jeden wiersz to jedna osoba. Wiersz z przykładem usuń.",
        "Wymagane: Nazwisko i Imię. Pozostałe kolumny można zostawić puste.",
        "PESEL: 11 cyfr. Data urodzenia i płeć uzupełnią się z PESEL.",
        "Daty: dd.mm.rrrr, np. 31.12.1990.",
        "Płeć: K albo M.",
        "Obywatelstwo, APT, Wydział, Stanowisko: nazwy jak w słownikach aplikacji "
        "(brakujące można dodać automatycznie przy imporcie).",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class RowResult:
    row: int
    status: str = "ok"  # ok | skip | error
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    values: dict[str, Any] = field(default_factory=dict)


__all__ = [
    "FIELDS",
    "FIELD_KEYS",
    "ImportFileError",
    "PeselError",
    "RowResult",
    "Table",
    "cell_text",
    "nationality_iso",
    "nice_name",
    "norm_header",
    "parse_date",
    "parse_pesel",
    "parse_sex",
    "read_table",
    "split_full_name",
    "suggest_mapping",
    "template_xlsx",
]
