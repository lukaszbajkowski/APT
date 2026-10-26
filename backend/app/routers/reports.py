"""Raporty i eksport do Excela. Zakres kierownika jak wszędzie; każdy eksport trafia do dziennika."""

import io
from collections import Counter, defaultdict
from datetime import date, datetime
from typing import Annotated, Any, Literal
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel
from sqlalchemy import func, select

from app import crypto, templating
from app.audit import audit
from app.deps import AuthContext, DbSession, require_roles
from app.models import Employee, EmployeeStatus, Role
from app.routers.employees import Filters, age_on, apply_scope, format_number, full_access

router = APIRouter(prefix="/api/reports", tags=["reports"])

ReportAuth = Annotated[AuthContext, Depends(require_roles(Role.admin, Role.hr, Role.manager))]
MAX_EXPORT = 20000
MAX_PDF_ROWS = 3000
ExportFormat = Annotated[Literal["xlsx", "pdf"], Query(alias="format")]
TZ = ZoneInfo("Europe/Warsaw")

NAVY = "1F3A73"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
MONTHS = ["sty", "lut", "mar", "kwi", "maj", "cze", "lip", "sie", "wrz", "paź", "lis", "gru"]
NONE = "(brak)"


def _sheet_table(ws: Worksheet, headers: list[str], rows: list[list[Any]], start_row: int = 1) -> None:
    """Nagłówek w kolorze aplikacji, filtr, zamrożony wiersz, szerokości kolumn."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if isinstance(v, date):
                cell.number_format = "DD.MM.YYYY"
    last = start_row + max(len(rows), 1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{last}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    # wydruk / PDF: poziomo, dopasowane do szerokości strony, nagłówek tabeli na każdej stronie
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{start_row}:{start_row}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.left.text = ws.title
    ws.oddFooter.right.text = "Strona &P z &N"
    for c, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(_fmt(r[c - 1])) for r in rows[:2000]]) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(max(width, 8), 50)


def _title(ws: Worksheet, title: str, user: Any, note: str | None = None) -> int:
    """Tytuł i podpis nad tabelą (wiersze 1–2/3); zwraca wiersz nagłówka tabeli."""
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Wygenerowano {datetime.now(TZ):%d.%m.%Y %H:%M} · {user.full_name}"
    ws["A2"].font = Font(italic=True, color="6B7280")
    if note:
        ws["A3"] = note
        ws["A3"].font = Font(color="6B7280")
        return 5
    return 4


def _fmt(v: Any) -> str:
    if isinstance(v, date):
        return "00.00.0000"
    return "" if v is None else str(v)


def _xlsx(wb: Workbook, name: str, fmt: str = "xlsx") -> Response:
    buf = io.BytesIO()
    wb.save(buf)
    if fmt == "pdf":
        try:
            pdf = templating.to_pdf(buf.getvalue(), suffix="xlsx")
        except templating.PdfUnavailable as exc:
            raise HTTPException(
                status.HTTP_503_SERVICE_UNAVAILABLE, detail=str(exc).replace("plik Word", "Excel")
            ) from exc
        pdf_name = name.rsplit(".", 1)[0] + ".pdf"
        return Response(
            pdf,
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{pdf_name}"',
                "Cache-Control": "no-store",
            },
        )
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"', "Cache-Control": "no-store"},
    )


def _stamp() -> str:
    return datetime.now(TZ).strftime("%Y-%m-%d_%H%M")


# --- eksport listy pracowników ----------------------------------------------------


@router.get("/employees.xlsx")
def export_employees(
    ctx: ReportAuth,
    request: Request,
    db: DbSession,
    filters: Filters,
    include_pesel: Annotated[bool, Query()] = False,
    fmt: ExportFormat = "xlsx",
) -> Response:
    user = ctx.user
    personal = full_access(user)
    if include_pesel and not personal:
        raise HTTPException(status.HTTP_403_FORBIDDEN, detail="Brak uprawnień do eksportu numerów PESEL.")
    emps = list(db.scalars(filters.query(user).order_by(*filters.order()).limit(MAX_EXPORT + 1)).unique())
    if len(emps) > MAX_EXPORT:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, detail=f"Za dużo wyników (ponad {MAX_EXPORT}). Zawęź filtry."
        )
    if fmt == "pdf" and len(emps) > MAX_PDF_ROWS:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail=f"PDF obejmuje najwyżej {MAX_PDF_ROWS} osób. Zawęź filtry albo pobierz Excel.",
        )

    today = date.today()
    headers = ["Nr", "Nazwisko", "Imię", "APT", "Wydział", "Stanowisko", "Obywatelstwo", "Płeć"]
    headers += ["Data urodzenia"] if personal else ["Wiek"]
    if include_pesel:
        headers.append("PESEL")
    headers += ["Data zatrudnienia", "Data zwolnienia", "Status", "Zdjęcie", "Uwagi"]

    rows = []
    for e in emps:
        row: list[Any] = [
            format_number(e.number),
            e.last_name,
            e.first_name,
            e.agency.name if e.agency else None,
            e.department.name if e.department else None,
            e.position.name if e.position else None,
            e.country.name if e.country else None,
            {"K": "kobieta", "M": "mężczyzna"}.get(e.sex.value) if e.sex else None,
            e.birth_date if personal else age_on(e.birth_date, today),
        ]
        if include_pesel:
            row.append(crypto.decrypt(e.pesel_enc) if e.pesel_enc else None)
        row += [
            e.hire_date,
            e.termination_date,
            "aktywny" if e.status == EmployeeStatus.active else "archiwum",
            "tak" if e.photo else "nie",
            e.notes,
        ]
        rows.append(row)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pracownicy"
    start = _title(ws, "Lista pracowników", user, f"Osób: {len(rows)}") if fmt == "pdf" else 1
    _sheet_table(ws, headers, rows, start_row=start)
    if include_pesel:
        ws.oddHeader.center.text = "Zawiera numery PESEL – dane poufne"

    if fmt == "pdf":
        # w PDF bez kolumny uwag – długi tekst rozpycha tabelę
        ws.delete_cols(len(headers))
    params = {k: v for k, v in request.query_params.multi_items() if k not in ("include_pesel", "format")}
    audit(
        db,
        request,
        "employees_exported",
        user=user,
        entity_type="report",
        details={"count": len(rows), "filters": params, "include_pesel": include_pesel, "format": fmt},
    )
    db.commit()
    return _xlsx(wb, f"pracownicy_{_stamp()}.xlsx", fmt)


# --- zestawienia ------------------------------------------------------------------


class Row(BaseModel):
    name: str
    count: int


class CrossRow(BaseModel):
    agency: str
    department: str
    count: int


class MonthRow(BaseModel):
    month: str  # RRRR-MM
    label: str
    hired: int
    terminated: int


class SummaryOut(BaseModel):
    generated_at: datetime
    active: int
    terminated: int
    by_agency: list[Row]
    by_department: list[Row]
    by_position: list[Row]
    by_country: list[Row]
    by_sex: list[Row]
    agency_department: list[CrossRow]
    months: list[MonthRow]
    missing_photo: int
    missing_pesel: int | None


def _rows(counter: Counter) -> list[Row]:
    return [Row(name=k, count=v) for k, v in sorted(counter.items(), key=lambda kv: (-kv[1], kv[0]))]


def _last_months(today: date, n: int = 12) -> list[tuple[int, int]]:
    out, y, m = [], today.year, today.month
    for _ in range(n):
        out.append((y, m))
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    return list(reversed(out))


def build_summary(db, user) -> SummaryOut:
    emps = list(db.scalars(apply_scope(select(Employee), user)).unique())
    active = [e for e in emps if e.status == EmployeeStatus.active]
    name = lambda r: r.name if r else NONE  # noqa: E731

    cross: dict[tuple[str, str], int] = defaultdict(int)
    for e in active:
        cross[(name(e.agency), name(e.department))] += 1

    today = date.today()
    months = []
    for y, m in _last_months(today):
        months.append(
            MonthRow(
                month=f"{y}-{m:02d}",
                label=f"{MONTHS[m - 1]} {y}",
                hired=sum(1 for e in emps if e.hire_date and (e.hire_date.year, e.hire_date.month) == (y, m)),
                terminated=sum(
                    1
                    for e in emps
                    if e.termination_date and (e.termination_date.year, e.termination_date.month) == (y, m)
                ),
            )
        )
    sex = Counter({"K": "kobiety", "M": "mężczyźni"}.get(e.sex.value if e.sex else "", NONE) for e in active)
    return SummaryOut(
        generated_at=datetime.now(TZ),
        active=len(active),
        terminated=len(emps) - len(active),
        by_agency=_rows(Counter(name(e.agency) for e in active)),
        by_department=_rows(Counter(name(e.department) for e in active)),
        by_position=_rows(Counter(name(e.position) for e in active)),
        by_country=_rows(Counter(name(e.country) for e in active)),
        by_sex=_rows(sex),
        agency_department=[
            CrossRow(agency=a, department=d, count=c)
            for (a, d), c in sorted(cross.items(), key=lambda kv: kv[0])
        ],
        months=months,
        missing_photo=sum(1 for e in active if e.photo is None),
        missing_pesel=sum(1 for e in active if e.pesel_hash is None) if full_access(user) else None,
    )


@router.get("/summary", response_model=SummaryOut)
def summary(ctx: ReportAuth, db: DbSession) -> SummaryOut:
    return build_summary(db, ctx.user)


@router.get("/summary.xlsx")
def summary_xlsx(ctx: ReportAuth, request: Request, db: DbSession, fmt: ExportFormat = "xlsx") -> Response:
    s = build_summary(db, ctx.user)
    wb = Workbook()

    ws = wb.active
    ws.title = "Podsumowanie"
    ws["A1"] = "Zestawienie pracowników"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Wygenerowano {s.generated_at:%d.%m.%Y %H:%M} · {ctx.user.full_name}"
    ws["A2"].font = Font(italic=True, color="6B7280")
    summary_rows: list[list[Any]] = [
        ["Aktywni pracownicy", s.active],
        ["W archiwum", s.terminated],
        ["Aktywni bez zdjęcia", s.missing_photo],
    ]
    if s.missing_pesel is not None:
        summary_rows.append(["Aktywni bez PESEL", s.missing_pesel])
    _sheet_table(ws, ["Wskaźnik", "Liczba"], summary_rows, start_row=4)
    ws.auto_filter.ref = None

    def counts(title: str, header: str, rows: list[Row]) -> None:
        sh = wb.create_sheet(title)
        total = sum(r.count for r in rows) or 1
        _sheet_table(
            sh, [header, "Aktywni", "Udział"], [[r.name, r.count, round(r.count / total, 4)] for r in rows]
        )
        for r in range(2, len(rows) + 2):
            sh.cell(row=r, column=3).number_format = "0.0%"

    counts("APT", "APT", s.by_agency)
    counts("Wydziały", "Wydział", s.by_department)
    cs = wb.create_sheet("APT × wydział")
    _sheet_table(
        cs, ["APT", "Wydział", "Aktywni"], [[r.agency, r.department, r.count] for r in s.agency_department]
    )
    counts("Stanowiska", "Stanowisko", s.by_position)
    counts("Obywatelstwo", "Obywatelstwo", s.by_country)
    counts("Płeć", "Płeć", s.by_sex)
    ms = wb.create_sheet("Ruch 12 mies.")
    _sheet_table(
        ms,
        ["Miesiąc", "Zatrudnieni", "Zwolnieni", "Saldo"],
        [[m.label, m.hired, m.terminated, m.hired - m.terminated] for m in s.months],
    )

    audit(
        db,
        request,
        "report_exported",
        user=ctx.user,
        entity_type="report",
        details={"report": "summary", "format": fmt},
    )
    db.commit()
    return _xlsx(wb, f"zestawienie_{_stamp()}.xlsx", fmt)


# --- raport szkoleń ---------------------------------------------------------------

TRAINING_STATUS_LABELS = {
    "valid": "ważne",
    "expiring": "kończy się",
    "expired": "wygasło",
    "none": "bez daty",
    "completed": "ukończone",
    "planned": "zaplanowane",
    "in_progress": "w trakcie",
    "overdue": "ponad standard",
}
TrainingStatusKey = Literal[
    "valid", "expiring", "expired", "none", "completed", "planned", "in_progress", "overdue"
]


class TrainingFilters:
    def __init__(
        self,
        type_id: int | None = None,
        program_id: int | None = None,
        department_id: int | None = None,
        mentor_id: int | None = None,
        status_: Annotated[list[TrainingStatusKey] | None, Query(alias="status")] = None,
        date_from: date | None = None,
        date_to: date | None = None,
        employees: Literal["active", "all"] = "active",
    ):
        self.type_id, self.program_id, self.department_id, self.mentor_id = (
            type_id,
            program_id,
            department_id,
            mentor_id,
        )
        self.status = status_ or []
        self.date_from, self.date_to, self.employees = date_from, date_to, employees


class TrainingRowOut(BaseModel):
    training_id: int
    employee_id: int
    number: str
    last_name: str
    first_name: str
    employee_department: str | None
    type_name: str
    name: str | None
    department: str | None
    mentor_name: str | None
    start_date: date | None
    exam_date: date | None
    end_date: date | None
    completed_on: date | None
    valid_until: date | None
    standard_days: int | None
    days: int | None
    status: str
    has_scan: bool
    notes: str | None


def build_trainings(db: Any, user: Any, f: TrainingFilters) -> list[TrainingRowOut]:
    from app.models import Training
    from app.routers.trainings import days_in_training, training_status

    q = select(Training, Employee).join(Employee, Employee.id == Training.employee_id)
    if f.employees == "active":
        q = q.where(Employee.status == EmployeeStatus.active)
    if f.type_id:
        q = q.where(Training.type_id == f.type_id)
    if f.program_id:
        q = q.where(Training.program_id == f.program_id)
    if f.mentor_id:
        q = q.where(Training.mentor_id == f.mentor_id)
    if f.department_id:
        # wydział szkolenia (stanowiskowe) albo wydział pracownika
        q = q.where(
            (Training.department_id == f.department_id)
            | ((Training.department_id.is_(None)) & (Employee.department_id == f.department_id))
        )
    # zakres dat: data ukończenia, a dla nieukończonych – data rozpoczęcia
    ref = func.coalesce(Training.completed_on, Training.start_date)
    if f.date_from:
        q = q.where(ref >= f.date_from)
    if f.date_to:
        q = q.where(ref <= f.date_to)
    q = apply_scope(q, user).order_by(Employee.last_name, Employee.first_name, ref.desc().nulls_first())
    today = date.today()
    out = []
    for tr, e in db.execute(q.limit(MAX_EXPORT + 1)).unique().all():
        st = training_status(tr, today)
        if f.status and st not in f.status:
            continue
        m = tr.mentor
        out.append(
            TrainingRowOut(
                training_id=tr.id,
                employee_id=e.id,
                number=format_number(e.number),
                last_name=e.last_name,
                first_name=e.first_name,
                employee_department=e.department.name if e.department else None,
                type_name=tr.type.name,
                name=tr.name,
                department=tr.department.name if tr.department else None,
                mentor_name=f"{m.last_name} {m.first_name}" if m else None,
                start_date=tr.start_date,
                exam_date=tr.exam_date,
                end_date=tr.end_date,
                completed_on=tr.completed_on,
                valid_until=tr.valid_until,
                standard_days=tr.standard_days,
                days=days_in_training(tr, today),
                status=st,
                has_scan=tr.scan_file is not None,
                notes=tr.notes,
            )
        )
    return out


TFilters = Annotated[TrainingFilters, Depends()]


class TrainingsReportOut(BaseModel):
    total: int
    by_status: dict[str, int]
    rows: list[TrainingRowOut]


@router.get("/trainings", response_model=TrainingsReportOut)
def trainings_report(
    ctx: ReportAuth, db: DbSession, filters: TFilters, limit: int = 200
) -> TrainingsReportOut:
    rows = build_trainings(db, ctx.user, filters)
    return TrainingsReportOut(
        total=len(rows),
        by_status=dict(Counter(r.status for r in rows)),
        rows=rows[: max(1, min(limit, 1000))],
    )


@router.get("/trainings.xlsx")
def trainings_export(
    ctx: ReportAuth, request: Request, db: DbSession, filters: TFilters, fmt: ExportFormat = "xlsx"
) -> Response:
    rows = build_trainings(db, ctx.user, filters)
    if len(rows) > MAX_EXPORT:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, detail=f"Za dużo wyników (ponad {MAX_EXPORT}). Zawęź filtry."
        )
    if fmt == "pdf" and len(rows) > MAX_PDF_ROWS:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail=f"PDF obejmuje najwyżej {MAX_PDF_ROWS} wpisów. Zawęź filtry albo pobierz Excel.",
        )
    headers = [
        "Nr",
        "Nazwisko",
        "Imię",
        "Wydział pracownika",
        "Rodzaj",
        "Szkolenie",
        "Wydział szkolenia",
        "Opiekun",
        "Rozpoczęcie",
        "Egzamin",
        "Zakończenie / ukończenie",
        "Ważne do",
        "Dni nauki",
        "Standard (dni)",
        "Status",
        "Skan",
    ]
    if fmt == "xlsx":
        headers.append("Uwagi")
    data = []
    for r in rows:
        row: list[Any] = [
            r.number,
            r.last_name,
            r.first_name,
            r.employee_department,
            r.type_name,
            r.name or r.type_name,
            r.department,
            r.mentor_name,
            r.start_date,
            r.exam_date,
            r.completed_on,
            r.valid_until,
            r.days,
            r.standard_days,
            TRAINING_STATUS_LABELS.get(r.status, r.status),
            "tak" if r.has_scan else "nie",
        ]
        if fmt == "xlsx":
            row.append(r.notes)
        data.append(row)
    wb = Workbook()
    ws = wb.active
    ws.title = "Szkolenia"
    by_status = Counter(TRAINING_STATUS_LABELS.get(r.status, r.status) for r in rows)
    note = f"Wpisów: {len(rows)}" + (
        "".join(f" · {k}: {v}" for k, v in by_status.most_common()) if rows else ""
    )
    _sheet_table(ws, headers, data, start_row=_title(ws, "Raport szkoleń", ctx.user, note))
    params = {k: v for k, v in request.query_params.multi_items() if k != "format"}
    audit(
        db,
        request,
        "report_exported",
        user=ctx.user,
        entity_type="report",
        details={"report": "trainings", "count": len(rows), "filters": params, "format": fmt},
    )
    db.commit()
    return _xlsx(wb, f"szkolenia_{_stamp()}.xlsx", fmt)
