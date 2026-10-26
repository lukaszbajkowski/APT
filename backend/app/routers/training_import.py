"""Import szkoleń z pliku (Administrator, Kadry) – jak import pracowników: analiza → podgląd → zapis.

Pracownika rozpoznajemy po numerze ewidencyjnym, PESEL albo nazwisku i imieniu (musi być jednoznaczne).
Rodzaj szkolenia – po nazwie ze słownika; jeśli kolumny rodzaju brak, a nazwa pasuje do pozycji katalogu
szkoleń stanowiskowych, przyjmujemy rodzaj tej pozycji.
"""

import io
import json
from dataclasses import dataclass, field
from datetime import date
from typing import Annotated, Any

from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import crypto
from app.audit import audit
from app.config import get_settings
from app.deps import DbSession, StaffAuth
from app.importer import (
    FieldSpec,
    ImportFileError,
    PeselError,
    Table,
    cell_text,
    norm_header,
    parse_date,
    parse_pesel,
    read_table,
    suggest_mapping,
)
from app.models import Department, Employee, Training, TrainingProgram, TrainingType
from app.routers.employees import format_number
from app.routers.trainings import TrainingIn, _apply

router = APIRouter(prefix="/api/trainings/import", tags=["trainings"])

FIELDS: tuple[FieldSpec, ...] = (
    FieldSpec(
        "number", "Nr ewidencyjny", ("nr", "numer", "nrewidencyjny", "numerewidencyjny", "nrpracownika")
    ),
    FieldSpec("pesel", "PESEL", ("pesel", "nrpesel", "numerpesel")),
    FieldSpec("last_name", "Nazwisko", ("nazwisko", "lastname", "surname")),
    FieldSpec("first_name", "Imię", ("imie", "imiona", "firstname")),
    FieldSpec(
        "full_name",
        "Nazwisko i imię (w jednej kolumnie)",
        ("nazwiskoiimie", "pracownik", "osoba", "fullname"),
    ),
    FieldSpec(
        "type", "Rodzaj szkolenia", ("rodzaj", "rodzajszkolenia", "typ", "typszkolenia", "uprawnienie")
    ),
    FieldSpec("name", "Nazwa szkolenia (stanowiskowe)", ("szkolenie", "nazwa", "nazwaszkolenia")),
    FieldSpec("department", "Wydział szkolenia", ("wydzial", "wydzialszkolenia", "dzial")),
    FieldSpec("start_date", "Data rozpoczęcia", ("datarozpoczecia", "rozpoczecie", "od", "start")),
    FieldSpec("exam_date", "Data egzaminu", ("dataegzaminu", "egzamin")),
    FieldSpec(
        "completed_on",
        "Data zakończenia / ukończenia",
        ("datazakonczenia", "zakonczenie", "dataukonczenia", "ukonczenie", "ukonczono", "do", "data"),
    ),
    FieldSpec("valid_until", "Ważne do", ("waznedo", "waznosc", "datawaznosci", "terminwaznosci")),
    FieldSpec("mentor", "Opiekun", ("opiekun", "prowadzacy", "instruktor", "mentor")),
    FieldSpec(
        "standard_days", "Czas nauki wg standardu (dni)", ("standard", "standarddni", "czasnauki", "dni")
    ),
    FieldSpec("notes", "Uwagi", ("uwagi", "notatki", "komentarz")),
)
FIELD_KEYS = {f.key for f in FIELDS}


class FieldOut(BaseModel):
    key: str
    label: str


class AnalyzeOut(BaseModel):
    file_name: str
    sheet: str | None
    headers: list[str]
    sample: list[list[str | None]]
    total_rows: int
    mapping: dict[str, int]


class RowOut(BaseModel):
    row: int
    status: str  # ok | skip | error
    errors: list[str]
    warnings: list[str]
    employee: str | None
    training: str | None
    dates: str | None


class Summary(BaseModel):
    total: int
    to_import: int
    skipped: int
    errors: int


class PreviewOut(BaseModel):
    summary: Summary
    rows: list[RowOut]


class CommitOut(BaseModel):
    imported: int
    skipped: int
    errors: int


@dataclass
class Result:
    row: int
    status: str = "ok"
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    employee: Employee | None = None
    body: TrainingIn | None = None
    label: str | None = None
    dates: str | None = None


def _k(v: str | None) -> str:
    return norm_header(v or "")


async def _load(file: UploadFile) -> tuple[Table, str]:
    limit = get_settings().photo_max_upload_mb * 1024 * 1024
    raw = await file.read(limit + 1)
    if len(raw) > limit:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Plik jest za duży.")
    name = file.filename or "plik"
    try:
        return read_table(raw, name), name
    except ImportFileError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


def _parse_mapping(raw: str, width: int) -> dict[str, int]:
    try:
        mapping = {str(k): int(v) for k, v in json.loads(raw).items() if v is not None}
    except (ValueError, TypeError, AttributeError) as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Nieprawidłowe przypisanie kolumn.") from exc
    if set(mapping) - FIELD_KEYS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Nieznane pola w przypisaniu.")
    if any(not 0 <= i < width for i in mapping.values()) or len(set(mapping.values())) != len(mapping):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Nieprawidłowe przypisanie kolumn.")
    who = {"number", "pesel", "full_name"} & set(mapping) or {"last_name", "first_name"} <= set(mapping)
    if not who:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="Przypisz kolumnę z pracownikiem: nr ewidencyjny, PESEL albo nazwisko i imię.",
        )
    if "type" not in mapping and "name" not in mapping:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="Przypisz kolumnę z rodzajem szkolenia (albo nazwą z katalogu).",
        )
    return mapping


class Lookup:
    """Słowniki i pracownicy wczytani raz na cały plik."""

    def __init__(self, db: Session):
        emps = list(db.scalars(select(Employee)).unique())
        self.by_number = {e.number: e for e in emps}
        self.by_hash = {e.pesel_hash: e for e in emps if e.pesel_hash}
        self.by_name: dict[str, list[Employee]] = {}
        for e in emps:
            self.by_name.setdefault(_k(f"{e.last_name} {e.first_name}"), []).append(e)
        self.types = {_k(t.name): t for t in db.scalars(select(TrainingType))}
        self.programs = list(db.scalars(select(TrainingProgram)).unique())
        self.departments = {_k(d.name): d for d in db.scalars(select(Department))}
        existing = db.execute(
            select(
                Training.employee_id,
                Training.type_id,
                Training.name,
                Training.completed_on,
                Training.start_date,
            )
        ).all()
        self.existing = {(r[0], r[1], _k(r[2]), r[3] or r[4]) for r in existing}

    def person(self, text: str) -> tuple[Employee | None, str | None]:
        """Opiekun / pracownik po numerze albo „Nazwisko Imię”."""
        digits = text.strip().lstrip("0")
        if digits.isdigit() and int(digits) in self.by_number:
            return self.by_number[int(digits)], None
        cands = self.by_name.get(_k(text), [])
        if not cands:
            parts = text.split()
            if len(parts) == 2:  # „Imię Nazwisko”
                cands = self.by_name.get(_k(f"{parts[1]} {parts[0]}"), [])
        if len(cands) == 1:
            return cands[0], None
        if len(cands) > 1:
            return None, f"„{text}” – kilka osób o tym nazwisku i imieniu, użyj nr ewidencyjnego"
        return None, f"nie znaleziono osoby „{text}”"


def _validate(db: Session, table: Table, mapping: dict[str, int], lk: Lookup) -> list[Result]:
    results: list[Result] = []
    seen: set[tuple] = set()
    for row, rownum in zip(table.rows, table.row_numbers, strict=True):

        def get(key: str, row: list[Any] = row) -> Any:
            i = mapping.get(key)
            return row[i] if i is not None and i < len(row) else None

        def text(key: str) -> str | None:
            return cell_text(get(key))

        r = Result(row=rownum)
        results.append(r)

        # --- pracownik
        e: Employee | None = None
        if text("number"):
            n = text("number").lstrip("0")
            e = lk.by_number.get(int(n)) if n.isdigit() else None
            if e is None:
                r.errors.append(f"nie ma pracownika o nr {text('number')}")
        elif text("pesel"):
            try:
                pesel = parse_pesel(get("pesel"))
                e = lk.by_hash.get(crypto.blind_index(pesel))
                if e is None:
                    r.errors.append("nie ma pracownika o tym numerze PESEL")
            except (PeselError, ValueError) as exc:
                r.errors.append(f"PESEL: {exc}")
        else:
            full = text("full_name") or " ".join(x for x in (text("last_name"), text("first_name")) if x)
            if full:
                e, err = lk.person(full)
                if err:
                    r.errors.append(err.replace("osoby", "pracownika"))
            else:
                r.errors.append("brak danych pracownika")
        r.employee = e

        # --- rodzaj / pozycja katalogu
        name = text("name")
        t = None
        if text("type"):
            t = lk.types.get(_k(text("type")))
            if t is None:
                r.errors.append(
                    f"nieznany rodzaj szkolenia „{text('type')}” – dodaj go w Szkolenia → Rodzaje"
                )
        elif name:
            progs = [p for p in lk.programs if _k(p.name) == _k(name)]
            if progs:
                t = db.get(TrainingType, progs[0].type_id)
            else:
                r.errors.append(f"nie podano rodzaju, a „{name}” nie ma w katalogu")

        dep = None
        if text("department"):
            dep = lk.departments.get(_k(text("department")))
            if dep is None:
                r.errors.append(f"nieznany wydział „{text('department')}”")

        program = None
        if t is not None and t.kind == "position" and name:
            progs = [p for p in lk.programs if p.type_id == t.id and _k(p.name) == _k(name)]
            program = next((p for p in progs if dep and p.department_id == dep.id), None) or (
                progs[0] if len(progs) == 1 else None
            )
            if program is None and progs:
                r.warnings.append("kilka pozycji katalogu o tej nazwie – bez powiązania z katalogiem")

        # --- daty i liczby
        dates: dict[str, date | None] = {}
        for key in ("start_date", "exam_date", "completed_on", "valid_until"):
            if get(key) is None or cell_text(get(key)) is None:
                dates[key] = None
                continue
            try:
                dates[key] = parse_date(get(key))
            except ValueError as exc:
                r.errors.append(f"{next(f.label for f in FIELDS if f.key == key)}: {exc}")
                dates[key] = None
        days = None
        if text("standard_days"):
            try:
                days = int(float(text("standard_days").replace(",", ".")))
            except ValueError:
                r.errors.append(f"czas nauki „{text('standard_days')}” nie jest liczbą")

        mentor = None
        if text("mentor"):
            mentor, err = lk.person(text("mentor"))
            if err:
                r.warnings.append(f"opiekun pominięty: {err}")

        if r.errors or e is None or t is None:
            r.status = "error"
            r.label = name or (t.name if t else text("type"))
            continue

        position = t.kind == "position"
        if position and program and days is None:
            days = program.standard_days
        if position and program and dep is None and program.department_id:
            dep = db.get(Department, program.department_id)
        if not position and name:
            r.warnings.append("nazwa pominięta – dotyczy tylko szkoleń stanowiskowych")
        body = TrainingIn(
            type_id=t.id,
            program_id=program.id if program else None,
            name=(program.name if program else name) if position else None,
            department_id=dep.id if dep and position else None,
            start_date=dates["start_date"],
            exam_date=dates["exam_date"],
            end_date=dates["completed_on"] if position else None,
            completed_on=dates["completed_on"],
            valid_until=dates["valid_until"],
            mentor_id=mentor.id if mentor and position and mentor.id != e.id else None,
            standard_days=days if position else None,
            notes=text("notes"),
        )
        # sprawdzenie reguł jak przy ręcznym wpisie (bez wymogu uprawnień opiekuna – to dane historyczne)
        probe = Training(employee_id=e.id)
        try:
            _apply(db, probe, body, check_mentor=False)
        except HTTPException as exc:
            r.status = "error"
            r.errors.append(str(exc.detail))
            continue
        r.body = body
        r.label = (probe.name or t.name) + (f" ({t.name})" if probe.name else "")
        ref = probe.completed_on or probe.start_date
        r.dates = (
            " – ".join(x.strftime("%d.%m.%Y") for x in (probe.start_date, probe.completed_on) if x) or None
        )
        if probe.valid_until:
            r.dates = (r.dates + " · " if r.dates else "") + f"ważne do {probe.valid_until:%d.%m.%Y}"
        key = (e.id, t.id, _k(probe.name), ref)
        if key in lk.existing:
            r.status = "skip"
            r.warnings.insert(0, "takie szkolenie już jest w kartotece")
        elif key in seen:
            r.status = "skip"
            r.warnings.insert(0, "powtórzony wiersz w pliku")
        seen.add(key)
    return results


def _row_out(r: Result) -> RowOut:
    e = r.employee
    return RowOut(
        row=r.row,
        status=r.status,
        errors=r.errors,
        warnings=r.warnings,
        employee=f"{e.last_name} {e.first_name} ({format_number(e.number)})" if e else None,
        training=r.label,
        dates=r.dates,
    )


def _summary(results: list[Result]) -> Summary:
    return Summary(
        total=len(results),
        to_import=sum(r.status == "ok" for r in results),
        skipped=sum(r.status == "skip" for r in results),
        errors=sum(r.status == "error" for r in results),
    )


@router.get("/fields", response_model=list[FieldOut])
def fields(ctx: StaffAuth) -> list[FieldOut]:
    return [FieldOut(key=f.key, label=f.label) for f in FIELDS]


@router.get("/template")
def template(ctx: StaffAuth) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = "Szkolenia"
    cols = ["Nr ewidencyjny", "Nazwisko", "Imię", "Rodzaj szkolenia", "Nazwa szkolenia", "Wydział szkolenia",
            "Data rozpoczęcia", "Data egzaminu", "Data zakończenia / ukończenia", "Ważne do", "Opiekun",
            "Czas nauki wg standardu (dni)", "Uwagi"]  # fmt: skip
    ws.append(cols)
    ws.append(
        ["000012", "Kowalski", "Jan", "Szkolenie BHP okresowe", "", "", "", "", "15.01.2026", "", "", "", ""]
    )
    ws.append(["000013", "Nowak", "Anna", "Szkolenie stanowiskowe", "Operator prasy", "Montaż", "01.03.2026",
               "10.03.2026", "12.03.2026", "", "Kowalski Jan", "10", ""])  # fmt: skip
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A73")
    for i, w in enumerate((14, 16, 12, 26, 22, 16, 15, 14, 22, 12, 18, 14, 30)):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=5000, min_col=1, max_col=1):
        for c in row:
            c.number_format = "@"
    info = wb.create_sheet("Instrukcja")
    for line in (
        "Jeden wiersz = jedno szkolenie jednej osoby. Wiersze z przykładami usuń.",
        "Pracownik: wystarczy nr ewidencyjny ALBO PESEL ALBO nazwisko i imię (gdy jednoznaczne).",
        "Rodzaj szkolenia: nazwa jak w aplikacji (Szkolenia → Rodzaje szkoleń).",
        "Szkolenia stanowiskowe: nazwa jak w katalogu – wydział i standard uzupełnią się z katalogu.",
        "Daty: dd.mm.rrrr. „Ważne do” puste = wyliczy się z okresu ważności rodzaju (jeśli ustawiony).",
        "Opiekun: nr ewidencyjny albo „Nazwisko Imię”.",
        "Szkolenia, które już są w kartotece (ta sama osoba, rodzaj, nazwa i data), zostaną pominięte.",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="import_szkolen_szablon.xlsx"'},
    )


@router.post("/analyze", response_model=AnalyzeOut)
async def analyze(ctx: StaffAuth, file: Annotated[UploadFile, File()]) -> AnalyzeOut:
    table, name = await _load(file)
    mapping = suggest_mapping(table.headers, FIELDS)
    if "last_name" in mapping and "first_name" in mapping:
        mapping.pop("full_name", None)
    return AnalyzeOut(
        file_name=name,
        sheet=table.sheet,
        headers=[h or f"(kolumna {i + 1})" for i, h in enumerate(table.headers)],
        sample=[[cell_text(c) for c in row] for row in table.rows[:5]],
        total_rows=len(table.rows),
        mapping=mapping,
    )


@router.post("/preview", response_model=PreviewOut)
async def preview(
    ctx: StaffAuth, db: DbSession, file: Annotated[UploadFile, File()], mapping: Annotated[str, Form()]
) -> PreviewOut:
    table, _ = await _load(file)
    results = _validate(db, table, _parse_mapping(mapping, len(table.headers)), Lookup(db))
    db.rollback()
    return PreviewOut(summary=_summary(results), rows=[_row_out(r) for r in results])


@router.post("/commit", response_model=CommitOut)
async def commit(
    ctx: StaffAuth,
    request: Request,
    db: DbSession,
    file: Annotated[UploadFile, File()],
    mapping: Annotated[str, Form()],
    skip_errors: Annotated[bool, Form()] = False,
) -> CommitOut:
    table, file_name = await _load(file)
    results = _validate(db, table, _parse_mapping(mapping, len(table.headers)), Lookup(db))
    summary = _summary(results)
    if summary.errors and not skip_errors:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail=f"W pliku jest {summary.errors} wierszy z błędami. Popraw je albo pomiń.",
        )
    ok = [r for r in results if r.status == "ok"]
    if not ok:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Brak wierszy do zaimportowania.")
    for r in ok:
        assert r.employee is not None and r.body is not None
        tr = Training(employee_id=r.employee.id, created_by_name=ctx.user.full_name)
        _apply(db, tr, r.body, check_mentor=False)
        db.add(tr)
    db.flush()
    audit(
        db,
        request,
        "trainings_imported",
        user=ctx.user,
        entity_type="import",
        details={
            "file": file_name[:200],
            "imported": len(ok),
            "skipped": summary.skipped,
            "errors_skipped": summary.errors,
        },
    )
    db.commit()
    return CommitOut(imported=len(ok), skipped=summary.skipped, errors=summary.errors)
