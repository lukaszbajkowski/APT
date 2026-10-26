import io
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.db import SessionLocal
from app.main import app
from app.models import Agency, AuditLog, Role, UserScope
from app.templating import soffice_path
from tests.conftest import login, make_pesel

MEDICOVER = (Path(__file__).parent / "fixtures" / "medicover_szablon.docx").read_bytes()
PESEL = make_pesel(date(1990, 1, 1), "M", 111)


def simple_docx(*paragraphs: str) -> bytes:
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def doc_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def all_text(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


@pytest.fixture
def admin(client, make_user):
    make_user("admin", Role.admin)
    login(client, "admin")
    return client


def upload(c, data: bytes, name="Testowy", filename="szablon.docx"):
    return c.post(
        "/api/templates",
        files={"file": (filename, data)},
        data={"name": name, "language": "pl"},
    )


def test_upload_medicover_detects_fields(admin):
    r = upload(admin, MEDICOVER, "Deklaracja Medicover")
    assert r.status_code == 201, r.text
    t = r.json()
    assert t["sensitive"] is True
    assert t["builtins"] == ["nazwisko_imie", "pesel", "data_urodzenia", "obywatelstwo", "plec"]
    keys = [f["key"] for f in t["fields"]]
    assert keys[:3] == ["deklaracja_typ", "program", "wariant"]
    typ = t["fields"][0]
    assert typ["type"] == "select"
    assert [o["value"] for o in typ["options"]] == ["przystapienie", "odstapienie", "zmiana"]
    assert typ["options"][0]["label"] == "przystąpienie"
    assert next(f for f in t["fields"] if f["key"] == "zam_miasto")["label"] == "Adres zamieszkania – miasto"


def test_bad_templates_rejected(admin):
    assert upload(admin, b"nie docx").status_code == 400
    r = upload(admin, simple_docx("Witaj {{ imie }"), "Zepsuty")
    assert r.status_code == 400
    assert "znacznikach" in r.json()["detail"]
    r = admin.post("/api/templates", files={"file": ("a.pdf", b"%PDF")}, data={"name": "x"})
    assert r.status_code == 400


def test_edit_fields(admin):
    t = upload(admin, simple_docx("{{ imie }} {{ miasto }} {{ data_od }}")).json()
    assert [(f["key"], f["type"]) for f in t["fields"]] == [("miasto", "text"), ("data_od", "date")]
    fields = [
        {"key": "data_od", "label": "Od kiedy", "type": "date", "required": True},
        {
            "key": "miasto",
            "label": "Miasto",
            "type": "select",
            "options": [{"value": "Wrocław", "label": "Wrocław"}],
        },
    ]
    r = admin.patch(f"/api/templates/{t['id']}", json={"fields": fields, "description": "Opis"})
    assert r.status_code == 200, r.text
    assert [f["label"] for f in r.json()["fields"]] == ["Od kiedy", "Miasto"]
    bad = admin.patch(f"/api/templates/{t['id']}", json={"fields": [{"key": "nieistnieje", "label": "X"}]})
    assert bad.status_code == 400
    bad = admin.patch(
        f"/api/templates/{t['id']}", json={"fields": [{"key": "miasto", "label": "X", "type": "select"}]}
    )
    assert bad.status_code == 400


def test_replace_file_keeps_field_settings(admin):
    t = upload(admin, simple_docx("{{ a }} {{ b }}")).json()
    admin.patch(
        f"/api/templates/{t['id']}",
        json={"fields": [{"key": "a", "label": "Pole A", "required": True}, {"key": "b", "label": "B"}]},
    )
    r = admin.put(
        f"/api/templates/{t['id']}/file", files={"file": ("v2.docx", simple_docx("{{ a }} {{ c }}"))}
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["version"] == 2
    assert [(f["key"], f["label"], f["required"]) for f in body["fields"]] == [
        ("a", "Pole A", True),
        ("c", "C", False),
    ]


def test_generate_document_fills_data(admin):
    with SessionLocal() as db:
        db.add(Agency(name="Agencja A"))
        db.commit()
    e = admin.post(
        "/api/employees",
        json={"first_name": "Jan", "last_name": "Kowalski", "pesel": PESEL, "country_id": 2, "agency_id": 1},
    ).json()
    t = upload(admin, MEDICOVER, "Deklaracja Medicover").json()
    values = {
        "deklaracja_typ": "przystapienie",
        "wariant": "rodzinny",
        "program": "Optimum",
        "zam_miasto": "Wrocław",
        "tel_kom": "600 100 200",
        "email": "jan@example.com",
    }
    r = admin.post(
        f"/api/employees/{e['id']}/documents", json={"template_id": t["id"], "values": values, "pdf": False}
    )
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["template_name"] == "Deklaracja Medicover"
    assert doc["contains_sensitive"] is True

    r = admin.get(f"/api/documents/{doc['id']}/download", params={"format": "docx"})
    assert r.status_code == 200
    assert 'filename="Deklaracja_Medicover_Kowalski_Jan.docx"' in r.headers["content-disposition"]
    text = all_text(r.content)
    assert "Kowalski Jan" in text
    assert PESEL in text
    assert "01.01.1990" in text
    assert "Ukraina" in text
    assert "Wrocław" in text
    assert "☒ rodzinnym" in text and "☐ pojedynczym" in text
    assert "☒ przystąpienia" in text
    assert "{{" not in text

    listed = admin.get(f"/api/employees/{e['id']}/documents").json()
    assert [d["id"] for d in listed] == [doc["id"]]
    with SessionLocal() as db:
        log = db.scalar(select(AuditLog).where(AuditLog.action == "document_generated"))
        assert log.details["contains_pesel"] is True


def test_form_validation(admin):
    e = admin.post("/api/employees", json={"first_name": "Jan", "last_name": "K"}).json()
    t = upload(admin, simple_docx("{{ typ }} {{ od }}", "{% if typ == 'a' %}A{% endif %}")).json()
    admin.patch(
        f"/api/templates/{t['id']}",
        json={
            "fields": [
                {
                    "key": "typ",
                    "label": "Typ",
                    "type": "select",
                    "required": True,
                    "options": [{"value": "a", "label": "A"}],
                },
                {"key": "od", "label": "Od", "type": "date"},
            ]
        },
    )
    url = f"/api/employees/{e['id']}/documents"
    r = admin.post(url, json={"template_id": t["id"], "values": {}, "pdf": False})
    assert r.status_code == 400 and "wymagane" in r.json()["detail"]
    r = admin.post(url, json={"template_id": t["id"], "values": {"typ": "zzz"}, "pdf": False})
    assert r.status_code == 400
    r = admin.post(
        url, json={"template_id": t["id"], "values": {"typ": "a", "od": "2026-13-01"}, "pdf": False}
    )
    assert r.status_code == 400
    r = admin.post(
        url, json={"template_id": t["id"], "values": {"typ": "a", "od": "2026-11-01"}, "pdf": False}
    )
    assert r.status_code == 201
    data = admin.get(f"/api/documents/{r.json()['id']}/download", params={"format": "docx"}).content
    assert doc_text(data).split("\n") == ["a 01.11.2026", "A"]


def _manager_setup(admin, make_user):
    with SessionLocal() as db:
        db.add(Agency(name="Agencja A"))
        db.commit()
    e = admin.post(
        "/api/employees", json={"first_name": "Jan", "last_name": "K", "agency_id": 1, "pesel": PESEL}
    ).json()
    mgr = make_user("kierownik", Role.manager)
    with SessionLocal() as db:
        db.add(UserScope(user_id=mgr.id, agency_id=1))
        db.commit()
    m = TestClient(app)
    login(m, "kierownik")
    return e, m


def test_manager_gets_sensitive_template_without_sensitive_values(admin, make_user):
    e, m = _manager_setup(admin, make_user)
    t = upload(
        admin, simple_docx("{{ nazwisko_imie }} PESEL:{{ pesel }} UR:{{ data_urodzenia }}"), "Z PESEL"
    ).json()
    r = m.post(f"/api/employees/{e['id']}/documents", json={"template_id": t["id"], "pdf": False})
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["masked"] is True
    assert doc["contains_sensitive"] is False
    # kierownik nie pobiera plików (tylko podgląd); treść sprawdzamy jako admin
    assert m.get(f"/api/documents/{doc['id']}/download", params={"format": "docx"}).status_code == 403
    data = admin.get(f"/api/documents/{doc['id']}/download", params={"format": "docx"}).content
    assert doc_text(data) == "K Jan PESEL: UR:"

    # Ten sam szablon u Kadr/Administratora – z PESEL; kierownik nie ma do niego dostępu.
    full = admin.post(
        f"/api/employees/{e['id']}/documents", json={"template_id": t["id"], "pdf": False}
    ).json()
    assert PESEL in doc_text(
        admin.get(f"/api/documents/{full['id']}/download", params={"format": "docx"}).content
    )
    assert m.get(f"/api/documents/{full['id']}/download", params={"format": "docx"}).status_code == 403
    assert m.get(f"/api/documents/{full['id']}/view").status_code == 403
    listed = {d["id"]: d for d in m.get(f"/api/employees/{e['id']}/documents").json()}
    assert listed[full["id"]]["accessible"] is False
    assert listed[full["id"]]["download"] == []
    assert listed[doc["id"]]["accessible"] is True
    assert listed[doc["id"]]["download"] == []


def test_allow_manager_switch(admin, make_user):
    e, m = _manager_setup(admin, make_user)
    t = upload(admin, simple_docx("{{ nazwisko_imie }}"), "Tylko kadry").json()
    admin.patch(f"/api/templates/{t['id']}", json={"allow_manager": False})
    assert m.get("/api/templates").json() == []
    r = m.post(f"/api/employees/{e['id']}/documents", json={"template_id": t["id"], "pdf": False})
    assert r.status_code == 403
    assert (
        m.post("/api/templates", files={"file": ("a.docx", simple_docx("x"))}, data={"name": "x"}).status_code
        == 403
    )


def test_download_formats_setting(admin):
    e = admin.post("/api/employees", json={"first_name": "Jan", "last_name": "K"}).json()
    t = upload(admin, simple_docx("{{ nazwisko_imie }}")).json()
    r = admin.patch(f"/api/templates/{t['id']}", json={"download_formats": "docx"})
    assert r.json()["download_formats"] == "docx"
    doc = admin.post(
        f"/api/employees/{e['id']}/documents", json={"template_id": t["id"], "pdf": False}
    ).json()
    assert doc["download"] == ["docx"]
    assert admin.get(f"/api/documents/{doc['id']}/download", params={"format": "pdf"}).status_code == 403
    assert admin.get(f"/api/documents/{doc['id']}/download", params={"format": "docx"}).status_code == 200
    assert admin.patch(f"/api/templates/{t['id']}", json={"download_formats": "xls"}).status_code == 422


def test_inactive_template_hidden_and_blocked(admin):
    e = admin.post("/api/employees", json={"first_name": "Jan", "last_name": "K"}).json()
    t = upload(admin, simple_docx("{{ imie }}")).json()
    admin.patch(f"/api/templates/{t['id']}", json={"is_active": False})
    assert admin.get("/api/templates", params={"active_only": True}).json() == []
    r = admin.post(f"/api/employees/{e['id']}/documents", json={"template_id": t["id"], "pdf": False})
    assert r.status_code == 400


@pytest.mark.skipif(soffice_path() is None, reason="brak LibreOffice")
def test_pdf_generation(admin):
    e = admin.post("/api/employees", json={"first_name": "Jan", "last_name": "Kowalski"}).json()
    t = upload(admin, simple_docx("Zaświadczenie dla {{ nazwisko_imie }}")).json()
    r = admin.post(f"/api/employees/{e['id']}/documents", json={"template_id": t["id"], "pdf": True})
    assert r.status_code == 201, r.text
    assert r.json()["has_pdf"] is True
    pdf = admin.get(f"/api/documents/{r.json()['id']}/download")
    assert pdf.content[:4] == b"%PDF"


@pytest.mark.skipif(soffice_path() is None, reason="brak LibreOffice")
def test_online_view_is_logged(admin):
    e = admin.post("/api/employees", json={"first_name": "Jan", "last_name": "Kowalski"}).json()
    t = upload(admin, simple_docx("Dla {{ nazwisko_imie }}")).json()
    admin.patch(f"/api/templates/{t['id']}", json={"download_formats": "docx"})
    doc = admin.post(
        f"/api/employees/{e['id']}/documents", json={"template_id": t["id"], "pdf": False}
    ).json()
    assert doc["has_pdf"] is True  # PDF powstaje do podglądu, nawet gdy do pobrania jest tylko Word
    assert doc["can_view"] is True
    r = admin.get(f"/api/documents/{doc['id']}/view")
    assert r.status_code == 200
    assert r.headers["content-disposition"].startswith("inline")
    assert r.content[:4] == b"%PDF"
    with SessionLocal() as db:
        assert db.scalar(select(AuditLog).where(AuditLog.action == "document_viewed")) is not None


def test_batch_zip_with_own_values(admin):
    import zipfile

    a = admin.post("/api/employees", json={"first_name": "Jan", "last_name": "Kowalski"}).json()
    b = admin.post("/api/employees", json={"first_name": "Ewa", "last_name": "Nowak"}).json()
    t = upload(admin, simple_docx("Skierowanie {{ nazwisko_imie }} na {{ badanie }}")).json()
    body = {
        "template_id": t["id"],
        "format": "docx",
        "items": [
            {"employee_id": a["id"], "values": {"badanie": "wstępne"}},
            {"employee_id": b["id"], "values": {"badanie": "okresowe"}},
        ],
    }
    r = admin.post("/api/documents/batch", json=body)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = sorted(z.namelist())
    assert len(names) == 2 and names[0].startswith(a["number_display"]) and names[0].endswith(".docx")
    texts = sorted(doc_text(z.read(n)) for n in names)
    assert texts == ["Skierowanie Kowalski Jan na wstępne", "Skierowanie Nowak Ewa na okresowe"]
    # każda osoba ma dokument w profilu
    assert len(admin.get(f"/api/employees/{b['id']}/documents").json()) == 1

    body["format"] = "pdf"
    t2 = admin.patch(f"/api/templates/{t['id']}", json={"download_formats": "docx"})
    assert t2.status_code == 200
    assert admin.post("/api/documents/batch", json=body).status_code == 400
    body["items"].append(body["items"][0])
    body["format"] = "docx"
    assert admin.post("/api/documents/batch", json=body).status_code == 400


def test_batch_validation_error_names_person_and_rolls_back(admin):
    a = admin.post("/api/employees", json={"first_name": "Jan", "last_name": "Kowalski"}).json()
    b = admin.post("/api/employees", json={"first_name": "Ewa", "last_name": "Nowak"}).json()
    t = upload(admin, simple_docx("{{ data_badania }}")).json()
    fields = t["fields"]
    fields[0]["type"] = "date"
    admin.patch(f"/api/templates/{t['id']}", json={"fields": fields})
    r = admin.post(
        "/api/documents/batch",
        json={
            "template_id": t["id"],
            "format": "docx",
            "items": [
                {"employee_id": a["id"], "values": {"data_badania": "2026-11-02"}},
                {"employee_id": b["id"], "values": {"data_badania": "jutro"}},
            ],
        },
    )
    assert r.status_code == 400 and "Nowak Ewa" in r.json()["detail"]
    assert admin.get(f"/api/employees/{a['id']}/documents").json() == []


def test_missing_pesel_entered_at_generation_is_saved(admin):
    from datetime import date as d

    from tests.conftest import make_pesel

    e = admin.post("/api/employees", json={"first_name": "Jan", "last_name": "Kowalski", "sex": "M"}).json()
    t = upload(admin, simple_docx("{{ nazwisko_imie }} PESEL {{ pesel }} ur. {{ data_urodzenia }}")).json()
    url = f"/api/employees/{e['id']}/documents"
    # zła suma kontrolna / niezgodna płeć
    assert admin.post(url, json={"template_id": t["id"], "pesel": "12345678901"}).status_code == 400
    woman = make_pesel(d(1990, 5, 1), "K")
    r = admin.post(url, json={"template_id": t["id"], "pesel": woman})
    assert r.status_code == 400 and "Płeć" in r.json()["detail"]
    assert admin.get(f"/api/employees/{e['id']}").json()["has_pesel"] is False

    good = make_pesel(d(1990, 5, 1), "M")
    r = admin.post(url, json={"template_id": t["id"], "pesel": good})
    assert r.status_code == 201, r.text
    emp = admin.get(f"/api/employees/{e['id']}").json()
    assert emp["has_pesel"] is True and emp["birth_date"] == "1990-05-01"
    text = doc_text(admin.get(f"/api/documents/{r.json()['id']}/download", params={"format": "docx"}).content)
    assert good in text and "01.05.1990" in text
    # duplikat u innej osoby
    other = admin.post("/api/employees", json={"first_name": "Ewa", "last_name": "Nowak"}).json()
    r = admin.post(f"/api/employees/{other['id']}/documents", json={"template_id": t["id"], "pesel": good})
    assert r.status_code == 409


def test_field_formats(admin):
    t = upload(admin, simple_docx("{{ zam_kod }} {{ tel_kom }} {{ uwagi }}")).json()
    fmts = {f["key"]: f["format"] for f in t["fields"]}
    assert fmts == {"zam_kod": "postal", "tel_kom": "phone", "uwagi": None}
    fields = t["fields"]
    fields[2]["format"] = "custom"
    assert admin.patch(f"/api/templates/{t['id']}", json={"fields": fields}).status_code == 422
    fields[2]["pattern"] = "AA/9999"
    assert admin.patch(f"/api/templates/{t['id']}", json={"fields": fields}).status_code == 200
    e = admin.post("/api/employees", json={"first_name": "Jan", "last_name": "Kowalski"}).json()
    url = f"/api/employees/{e['id']}/documents"
    bad = admin.post(
        url, json={"template_id": t["id"], "values": {"zam_kod": "5000", "tel_kom": "600100200"}}
    )
    assert bad.status_code == 400 and "00-000" in bad.json()["detail"]
    ok = admin.post(
        url,
        json={
            "template_id": t["id"],
            "values": {"zam_kod": "50001", "tel_kom": "600100200", "uwagi": "ab1234"},
        },
    )
    assert ok.status_code == 201, ok.text
    text = doc_text(
        admin.get(f"/api/documents/{ok.json()['id']}/download", params={"format": "docx"}).content
    )
    assert text == "50-001 600 100 200 AB/1234"
    assert {f["value"] for f in admin.get("/api/templates/formats").json()} >= {"postal", "nip", "iban"}
