import io
import json
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.importer import ImportFileError, read_table, suggest_mapping
from app.models import Agency, AuditLog, Department, Employee, Position, Role
from tests.conftest import login, make_pesel

P1 = make_pesel(date(1990, 1, 1), "M", 101)
P2 = make_pesel(date(2003, 5, 7), "K", 202)  # zaczyna się od 0 – Excel go „zjada”


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADERS = [
    "Lp.",
    "Nazwisko i imię",
    "PESEL",
    "Obywatelstwo",
    "Agencja",
    "Dział",
    "Stanowisko",
    "Zatrudniony od",
]


@pytest.fixture
def hr(client, make_user):
    make_user("kadry", Role.hr)
    login(client, "kadry")
    with SessionLocal() as db:
        d = Department(name="Montaż")
        db.add_all([Agency(name="Agencja A", code="AA"), d])
        db.flush()
        db.add(Position(name="Operator", department_id=d.id))
        db.commit()
    return client


def post(c, path, data: bytes, name="plik.xlsx", **form):
    form = {
        k: (json.dumps(v) if isinstance(v, dict) else str(v).lower() if isinstance(v, bool) else v)
        for k, v in form.items()
    }
    return c.post(f"/api/import/{path}", files={"file": (name, data)}, data=form)


# ---------- moduł importer ----------


def test_suggest_mapping_polish_headers():
    m = suggest_mapping(HEADERS + ["Data urodzenia (dd.mm.rrrr)", "Płeć"])
    assert m == {
        "full_name": 1,
        "pesel": 2,
        "country": 3,
        "agency": 4,
        "department": 5,
        "position": 6,
        "hire_date": 7,
        "birth_date": 8,
        "sex": 9,
    }


def test_read_table_skips_title_rows_and_blank_lines():
    data = xlsx(
        [["Lista pracowników – październik"], [], ["Nazwisko", "Imię"], ["A", "B"], [None, None], ["C", "D"]]
    )
    t = read_table(data, "x.xlsx")
    assert t.headers == ["Nazwisko", "Imię"]
    assert t.rows == [["A", "B"], ["C", "D"]]
    assert t.row_numbers == [4, 6]


def test_read_csv_semicolon_cp1250():
    raw = "Nazwisko;Imię;Płeć\nŻółć;Łucja;K\n".encode("cp1250")
    t = read_table(raw, "x.csv")
    assert t.rows == [["Żółć", "Łucja", "K"]]


def test_unsupported_formats():
    with pytest.raises(ImportFileError, match="xls"):
        read_table(b"\xd0\xcf", "stary.xls")
    with pytest.raises(ImportFileError):
        read_table(b"abc", "x.pdf")


# ---------- API ----------


def test_template_downloads(hr):
    r = hr.get("/api/import/template")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb["Pracownicy"]["A1"].value == "Nazwisko"
    assert "Instrukcja" in wb.sheetnames


def test_analyze_suggests_mapping(hr):
    data = xlsx(
        [HEADERS, [1, "KOWALSKI JAN", P1, "ukraińskie", "AA", "Montaż", "Operator", datetime(2026, 3, 1)]]
    )
    r = post(hr, "analyze", data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total_rows"] == 1
    assert body["mapping"]["full_name"] == 1
    assert body["sample"][0][7] == "01.03.2026"


def _file():
    return xlsx(
        [
            HEADERS,
            [1, "KOWALSKI JAN", P1, "ukraińskie", "AA", "Montaż", "Operator", datetime(2026, 3, 1)],
            [2, "Nowak Anna", int(P2), "Polska", "Agencja Nowa", "Magazyn", "Magazynier", "15.02.2026"],
            [3, "Zły Pesel", "12345678901", "PL", "AA", None, None, None],
            [4, "Kowalczyk", None, None, None, None, None, None],
            [5, "Kowalski Jan", P1, None, None, None, None, None],  # ten sam PESEL co wiersz 2
        ]
    )


MAPPING = {
    "full_name": 1,
    "pesel": 2,
    "country": 3,
    "agency": 4,
    "department": 5,
    "position": 6,
    "hire_date": 7,
}


def test_preview_reports_each_row(hr):
    r = post(hr, "preview", _file(), mapping=MAPPING)
    assert r.status_code == 200, r.text
    body = r.json()
    rows = {x["row"]: x for x in body["rows"]}
    assert rows[2]["status"] == "ok"
    assert rows[2]["name"] == "Kowalski Jan"
    assert rows[2]["country"] == "Ukraina"
    assert rows[2]["agency"] == "Agencja A"  # dopasowane po skrócie
    assert rows[2]["birth_date"] == "1990-01-01"
    assert rows[2]["pesel_masked"].endswith(P1[-2:])
    assert P1 not in r.text  # pełny PESEL nie wraca do przeglądarki

    assert rows[3]["status"] == "error"  # nowa APT / wydział / stanowisko bez zgody na dodanie
    assert any("Agencja Nowa" in e for e in rows[3]["errors"])
    assert rows[4]["status"] == "error" and "PESEL" in rows[4]["errors"][0]
    assert rows[5]["status"] == "error"
    assert rows[6]["status"] == "error" and "wierszu 2" in rows[6]["errors"][0]
    assert body["summary"] == {
        "total": 5,
        "to_import": 1,
        "skipped": 0,
        "errors": 4,
        "new_dictionary_items": {},
    }

    # Z dodawaniem brakujących słowników wiersz 3 jest poprawny, a PESEL z zerem odtworzony.
    r = post(hr, "preview", _file(), mapping=MAPPING, create_missing=True)
    body = r.json()
    rows = {x["row"]: x for x in body["rows"]}
    assert rows[3]["status"] == "ok", rows[3]
    assert rows[3]["birth_date"] == "2003-05-07"
    assert body["summary"]["new_dictionary_items"] == {
        "agencies": ["Agencja Nowa"],
        "departments": ["Magazyn"],
        "positions": ["Magazynier"],
    }


def test_commit_requires_decision_on_errors(hr):
    r = post(hr, "commit", _file(), mapping=MAPPING, create_missing=True)
    assert r.status_code == 400
    with SessionLocal() as db:
        assert db.scalar(select(Employee).limit(1)) is None


def test_commit_imports_valid_rows_in_one_go(hr):
    r = post(hr, "commit", _file(), mapping=MAPPING, create_missing=True, skip_errors=True)
    assert r.status_code == 200, r.text
    assert r.json() == {
        "imported": 2,
        "skipped": 0,
        "errors": 3,
        "first_number": "000001",
        "last_number": "000002",
    }

    with SessionLocal() as db:
        emps = {e.last_name: e for e in db.scalars(select(Employee))}
        assert set(emps) == {"Kowalski", "Nowak"}
        assert emps["Nowak"].department.name == "Magazyn"
        assert emps["Nowak"].position.department_id == emps["Nowak"].department_id
        assert emps["Kowalski"].position.name == "Operator"
        assert emps["Kowalski"].hire_date == date(2026, 3, 1)
        assert emps["Kowalski"].sex.value == "M"
        assert db.scalar(select(Agency).where(Agency.name == "Agencja Nowa")) is not None
        actions = list(db.scalars(select(AuditLog.action)))
        assert actions.count("employee_created") == 2
        assert actions.count("import_completed") == 1

    # Ten sam plik drugi raz: osoby z PESEL już są -> pominięte.
    r = post(hr, "preview", _file(), mapping=MAPPING, create_missing=True)
    s = r.json()["summary"]
    assert s["to_import"] == 0
    assert s["skipped"] == 3  # wiersze 2, 3 i powtórzenie z wiersza 6

    # PESEL z importu jest odczytywalny
    kid = emps["Nowak"].id
    assert hr.post(f"/api/employees/{kid}/pesel").json()["pesel"] == P2


def test_separate_name_columns_and_possible_duplicate(hr):
    hr.post("/api/employees", json={"first_name": "Ola", "last_name": "Lis", "birth_date": "1999-09-09"})
    data = xlsx(
        [
            ["Nazwisko", "Imię", "Data urodzenia", "Płeć"],
            ["Lis", "Ola", "09.09.1999", "K"],
            ["Wilk", "Adam", "1985-02-03", "mężczyzna"],
            ["Sowa", "Ewa", "31.02.1990", "x"],
        ]
    )
    r = post(hr, "preview", data, mapping={"last_name": 0, "first_name": 1, "birth_date": 2, "sex": 3})
    rows = {x["row"]: x for x in r.json()["rows"]}
    assert rows[2]["status"] == "skip"
    assert rows[3]["status"] == "ok"
    assert rows[4]["status"] == "error"
    assert len(rows[4]["errors"]) == 2


def test_mapping_validation(hr):
    data = xlsx([["Nazwisko", "Imię"], ["A", "B"]])
    assert post(hr, "preview", data, mapping={"last_name": 0}).status_code == 400
    assert post(hr, "preview", data, mapping={"last_name": 0, "first_name": 0}).status_code == 400
    assert post(hr, "preview", data, mapping={"last_name": 0, "first_name": 5}).status_code == 400
    assert post(hr, "preview", data, mapping={"hacker": 0, "first_name": 1}).status_code == 400


def test_only_staff_can_import(client, make_user):
    make_user("kierownik", Role.manager)
    login(client, "kierownik")
    assert client.get("/api/import/template").status_code == 403
    assert post(client, "analyze", xlsx([["A", "B"], [1, 2]])).status_code == 403
