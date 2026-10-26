import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AuditLog, Department
from app.templating import soffice_path
from tests.test_trainings import _emp, _types


def _setup(c):
    with SessionLocal() as db:
        db.add_all([Department(name="Montaż"), Department(name="Magazyn")])
        db.commit()
        deps = {d.name: d.id for d in db.scalars(select(Department))}
    types = _types(c)
    pos, bhp = types["Szkolenie stanowiskowe"]["id"], types["Szkolenie BHP okresowe"]["id"]
    today = date.today()
    a = _emp(c, department_id=deps["Magazyn"])
    b = _emp(c, "Jan", "Kowal", department_id=deps["Magazyn"])
    c.post(f"/api/employees/{a['id']}/trainings", json={"type_id": bhp, "completed_on": "2026-01-10"})
    body = {
        "type_id": pos,
        "name": "Operator prasy",
        "department_id": deps["Montaż"],
        "start_date": (today - timedelta(days=12)).isoformat(),
        "standard_days": 10,
    }
    c.post(f"/api/employees/{b['id']}/trainings", json=body)
    return deps, pos, bhp


def test_trainings_report_filters(admin_client):
    c = admin_client
    deps, pos, bhp = _setup(c)
    r = c.get("/api/reports/trainings").json()
    assert r["total"] == 2 and r["by_status"] == {"valid": 1, "overdue": 1}
    assert [
        x["last_name"] for x in c.get("/api/reports/trainings", params={"status": "overdue"}).json()["rows"]
    ] == ["Kowal"]
    # wydział szkolenia ma pierwszeństwo przed wydziałem pracownika
    by_dep = c.get("/api/reports/trainings", params={"department_id": deps["Montaż"]}).json()
    assert [x["name"] for x in by_dep["rows"]] == ["Operator prasy"]
    by_dep = c.get("/api/reports/trainings", params={"department_id": deps["Magazyn"]}).json()
    assert [x["type_name"] for x in by_dep["rows"]] == ["Szkolenie BHP okresowe"]
    assert (
        c.get("/api/reports/trainings", params={"type_id": bhp, "date_from": "2026-02-01"}).json()["total"]
        == 0
    )


def test_trainings_export_xlsx_and_audit(admin_client):
    c = admin_client
    _setup(c)
    r = c.get("/api/reports/trainings.xlsx", params={"status": ["valid", "overdue"]})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = [[x.value for x in row] for row in ws.iter_rows(min_row=5)]
    assert ws["A1"].value == "Raport szkoleń" and "Wpisów: 2" in ws["A3"].value
    assert rows[0][0] == "Nr" and "Uwagi" in rows[0] and len(rows) == 3
    assert {x[rows[0].index("Status")] for x in rows[1:]} == {"ważne", "ponad standard"}
    with SessionLocal() as db:
        log = db.scalar(select(AuditLog).where(AuditLog.action == "report_exported"))
        assert log.details["report"] == "trainings" and log.details["count"] == 2


@pytest.mark.skipif(soffice_path() is None, reason="brak LibreOffice")
def test_reports_pdf(admin_client):
    c = admin_client
    _setup(c)
    for url in ("/api/reports/trainings.xlsx", "/api/reports/summary.xlsx", "/api/reports/employees.xlsx"):
        r = c.get(url, params={"format": "pdf"})
        assert r.status_code == 200, (url, r.text)
        assert r.content.startswith(b"%PDF") and r.headers["content-type"] == "application/pdf"
        assert ".pdf" in r.headers["content-disposition"]
