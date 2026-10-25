import io
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.main import app
from app.models import Agency, AuditLog, Department, Role, UserScope
from tests.conftest import login, make_pesel


def _setup(c):
    with SessionLocal() as db:
        db.add_all([Agency(name="A"), Agency(name="B"), Department(name="Magazyn")])
        db.commit()
    p = make_pesel(date(1990, 5, 1), "K")
    c.post(
        "/api/employees",
        json={
            "first_name": "Ola",
            "last_name": "Nowak",
            "agency_id": 1,
            "department_id": 1,
            "pesel": p,
            "hire_date": "2026-03-01",
        },
    )
    c.post(
        "/api/employees",
        json={"first_name": "Jan", "last_name": "Kowal", "agency_id": 2, "hire_date": "2026-04-01"},
    )
    e = c.post("/api/employees", json={"first_name": "Ewa", "last_name": "Stara", "agency_id": 1}).json()
    c.post(f"/api/employees/{e['id']}/terminate", json={"termination_date": date.today().isoformat()})
    return p


def _rows(r):
    ws = load_workbook(io.BytesIO(r.content)).active
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_export_with_filters_and_pesel(admin_client):
    c = admin_client
    pesel = _setup(c)
    r = c.get("/api/reports/employees.xlsx", params={"agency_id": 1})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    rows = _rows(r)
    assert rows[0][:3] == ["Nr", "Nazwisko", "Imię"] and "PESEL" not in rows[0]
    assert [x[1] for x in rows[1:]] == ["Nowak"]  # domyślnie aktywni z APT A
    assert isinstance(rows[1][rows[0].index("Data urodzenia")], object)

    rows = _rows(c.get("/api/reports/employees.xlsx", params={"status": "all", "include_pesel": "true"}))
    assert pesel in [x[rows[0].index("PESEL")] for x in rows[1:]]
    assert len(rows) == 4
    with SessionLocal() as db:
        logs = list(db.scalars(select(AuditLog).where(AuditLog.action == "employees_exported")))
        assert [lg.details["include_pesel"] for lg in logs] == [False, True]
        assert logs[0].details["filters"]["agency_id"] == "1"


def test_manager_scope_no_pesel(client, make_user):
    make_user("kadry", Role.hr)
    login(client, "kadry")
    _setup(client)
    mgr = make_user("kier", Role.manager)
    with SessionLocal() as db:
        db.add(UserScope(user_id=mgr.id, agency_id=2))
        db.commit()
    m = TestClient(app)
    login(m, "kier")
    rows = _rows(m.get("/api/reports/employees.xlsx"))
    assert [x[1] for x in rows[1:]] == ["Kowal"] and "Wiek" in rows[0] and "Data urodzenia" not in rows[0]
    assert m.get("/api/reports/employees.xlsx", params={"include_pesel": "true"}).status_code == 403
    s = m.get("/api/reports/summary").json()
    assert s["active"] == 1 and s["missing_pesel"] is None
    make_user("podglad", Role.viewer)
    v = TestClient(app)
    login(v, "podglad")
    assert v.get("/api/reports/summary").status_code == 403


def test_summary_and_xlsx(admin_client):
    c = admin_client
    _setup(c)
    s = c.get("/api/reports/summary").json()
    assert (s["active"], s["terminated"]) == (2, 1)
    assert {r["name"]: r["count"] for r in s["by_agency"]} == {"A": 1, "B": 1}
    assert {(r["agency"], r["department"]) for r in s["agency_department"]} == {
        ("A", "Magazyn"),
        ("B", "(brak)"),
    }
    assert len(s["months"]) == 12 and s["months"][-1]["terminated"] == 1
    assert s["missing_pesel"] == 1
    r = c.get("/api/reports/summary.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == [
        "Podsumowanie",
        "APT",
        "Wydziały",
        "APT × wydział",
        "Stanowiska",
        "Obywatelstwo",
        "Płeć",
        "Ruch 12 mies.",
    ]
    assert wb["APT"]["A1"].value == "APT"
